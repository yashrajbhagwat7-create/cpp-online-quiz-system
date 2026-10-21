#!/usr/bin/env python3
"""
Simple, Student-Friendly PowerPoint Presentation Generator
Project: Online Examination & Quiz System (C++ Mini Project)
Target: 14 clean, academic slides without unnecessary business jargon.
"""

import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# ==============================================================================
# COLOR PALETTE (Clean Academic Theme)
# ==============================================================================
C_NAVY_DARK    = RGBColor(30, 41, 59)      # Slate 800 (#1E293B)
C_SLATE_900    = RGBColor(15, 23, 42)      # Slate 900 (#0F172A)
C_SLATE_700    = RGBColor(51, 65, 85)      # Slate 700 (#334155)
C_SLATE_600    = RGBColor(71, 85, 105)     # Slate 600 (#475569)
C_SLATE_400    = RGBColor(148, 163, 184)   # Slate 400 (#94A3B8)
C_SLATE_200    = RGBColor(226, 232, 240)   # Slate 200 (#E2E8F0)
C_SLATE_100    = RGBColor(241, 245, 249)   # Slate 100 (#F1F5F9)
C_BG_LIGHT     = RGBColor(248, 250, 252)   # Slate 50  (#F8FAFC)
C_WHITE        = RGBColor(255, 255, 255)

C_BLUE_PRIMARY = RGBColor(37, 99, 235)     # Blue 600 (#2563EB)
C_BLUE_ACCENT  = RGBColor(14, 165, 233)    # Sky 500 (#0EA5E9)
C_BLUE_LIGHT   = RGBColor(239, 246, 255)   # Blue 50 (#EFF6FF)
C_TEAL         = RGBColor(16, 185, 129)    # Emerald 500 (#10B981)
C_TEAL_LIGHT   = RGBColor(236, 253, 245)   # Emerald 50 (#ECFDF5)
C_PURPLE       = RGBColor(124, 58, 237)    # Violet 600 (#7C3AED)
C_PURPLE_LIGHT = RGBColor(245, 243, 255)   # Violet 50 (#F5F3FF)
C_AMBER        = RGBColor(217, 119, 6)     # Amber 600 (#D97706)

FONT_MAIN = "Segoe UI"
FONT_CODE = "Consolas"
TOTAL_SLIDES = 14

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # --- Helper: Add Card Rectangle ---
    def add_card(slide, left, top, width, height, bg_color=C_WHITE, border_color=C_SLATE_200, border_width=1):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(border_width)
        else:
            shape.line.fill.background()
        return shape

    # --- Helper: Add Slide Header ---
    def add_header(slide, tag_text, title_text, subtitle_text=None, is_dark=False):
        # Tag category
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.3))
        tf_tag = tag_box.text_frame
        tf_tag.word_wrap = True
        tf_tag.margin_left = tf_tag.margin_top = tf_tag.margin_right = tf_tag.margin_bottom = 0
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = tag_text.upper()
        p_tag.font.name = FONT_MAIN
        p_tag.font.size = Pt(10)
        p_tag.font.bold = True
        p_tag.font.color.rgb = C_BLUE_ACCENT if is_dark else C_BLUE_PRIMARY

        # Title & Subtitle in single text box
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.7), Inches(0.95))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_title = tf.paragraphs[0]
        p_title.text = title_text
        p_title.font.name = FONT_MAIN
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = C_WHITE if is_dark else C_SLATE_900

        if subtitle_text:
            p_sub = tf.add_paragraph()
            p_sub.text = subtitle_text
            p_sub.font.name = FONT_MAIN
            p_sub.font.size = Pt(12)
            p_sub.font.color.rgb = C_SLATE_400 if is_dark else C_SLATE_600
            p_sub.space_before = Pt(3)

    # --- Helper: Add Slide Footer ---
    def add_footer(slide, current_slide, is_dark=False):
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(9.0), Inches(0.3))
        tf = footer_box.text_frame
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = "C++ Object-Oriented Online Examination & Quiz System | Lab Mini Project"
        p.font.name = FONT_MAIN
        p.font.size = Pt(9)
        p.font.color.rgb = C_SLATE_400 if is_dark else C_SLATE_600

        num_box = slide.shapes.add_textbox(Inches(10.5), Inches(7.05), Inches(2.0), Inches(0.3))
        tf_num = num_box.text_frame
        tf_num.margin_left = tf_num.margin_top = tf_num.margin_right = tf_num.margin_bottom = 0
        p_num = tf_num.paragraphs[0]
        p_num.alignment = PP_ALIGN.RIGHT
        p_num.text = f"{current_slide} / {TOTAL_SLIDES}"
        p_num.font.name = FONT_MAIN
        p_num.font.size = Pt(9)
        p_num.font.bold = True
        p_num.font.color.rgb = C_SLATE_400 if is_dark else C_SLATE_600

    # --- Helper: Add Bullet Item ---
    def add_bullet_item(tf, prefix, text, font_size=11, space_before=7, bullet_color=C_BLUE_PRIMARY):
        p = tf.add_paragraph()
        p.font.name = FONT_MAIN
        p.font.size = Pt(font_size)
        p.space_before = Pt(space_before)
        p.line_spacing = 1.15

        r_bullet = p.add_run()
        r_bullet.text = "•  "
        r_bullet.font.bold = True
        r_bullet.font.size = Pt(font_size + 1)
        r_bullet.font.color.rgb = bullet_color

        if prefix:
            r_pre = p.add_run()
            r_pre.text = prefix + ": "
            r_pre.font.bold = True
            r_pre.font.color.rgb = C_SLATE_900

        r_txt = p.add_run()
        r_txt.text = text
        r_txt.font.color.rgb = C_SLATE_700

    # =========================================================================
    # SLIDE 1: TITLE SLIDE (Clean Modern Tech)
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    add_card(slide1, 0, 0, Inches(13.333), Inches(7.5), bg_color=C_SLATE_900, border_color=None)
    add_card(slide1, 0, 0, Inches(13.333), Inches(0.1), bg_color=C_BLUE_PRIMARY, border_color=None)

    # Top Tag
    tb_t1 = slide1.shapes.add_textbox(Inches(1.0), Inches(1.3), Inches(11.3), Inches(0.4))
    tf_t1 = tb_t1.text_frame
    p_t1 = tf_t1.paragraphs[0]
    p_t1.text = "OBJECT-ORIENTED PROGRAMMING (C++) LAB MINI PROJECT"
    p_t1.font.name = FONT_MAIN
    p_t1.font.size = Pt(11)
    p_t1.font.bold = True
    p_t1.font.color.rgb = C_BLUE_ACCENT

    # Main Project Title
    tb_title = slide1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(1.8))
    tf_title = tb_title.text_frame
    tf_title.word_wrap = True
    pt1 = tf_title.paragraphs[0]
    pt1.text = "Online Examination & Quiz System"
    pt1.font.name = FONT_MAIN
    pt1.font.size = Pt(38)
    pt1.font.bold = True
    pt1.font.color.rgb = C_WHITE

    pt2 = tf_title.add_paragraph()
    pt2.text = "A File-Driven Console Assessment Engine Built with Modern C++"
    pt2.font.name = FONT_MAIN
    pt2.font.size = Pt(16)
    pt2.font.color.rgb = C_SLATE_400
    pt2.space_before = Pt(8)

    # 3 Simple Feature Badges
    badges = [
        ("Core OOP Principles", "Abstraction, Inheritance, Polymorphism"),
        ("Dynamic File I/O", "questions.txt & results.txt"),
        ("Single-File Clean Build", "Easy lab compilation & viva review")
    ]
    for i, (b_title, b_sub) in enumerate(badges):
        bx = Inches(1.0 + i * 3.85)
        by = Inches(3.8)
        bw = Inches(3.6)
        add_card(slide1, bx, by, bw, Inches(1.05), bg_color=C_NAVY_DARK, border_color=C_SLATE_700)
        tb_b = slide1.shapes.add_textbox(bx, by, bw, Inches(1.05))
        tf_b = tb_b.text_frame
        tf_b.margin_left = tf_b.margin_right = Inches(0.2)
        tf_b.margin_top = Inches(0.18)
        pb1 = tf_b.paragraphs[0]
        pb1.text = b_title
        pb1.font.name = FONT_MAIN
        pb1.font.size = Pt(12)
        pb1.font.bold = True
        pb1.font.color.rgb = C_WHITE
        pb2 = tf_b.add_paragraph()
        pb2.text = b_sub
        pb2.font.name = FONT_MAIN
        pb2.font.size = Pt(10)
        pb2.font.color.rgb = C_SLATE_400
        pb2.space_before = Pt(3)

    # Student Submission Card
    meta_box = add_card(slide1, Inches(1.0), Inches(5.2), Inches(11.33), Inches(1.4), bg_color=C_NAVY_DARK, border_color=C_BLUE_PRIMARY, border_width=1.5)
    tb_meta = slide1.shapes.add_textbox(Inches(1.2), Inches(5.35), Inches(10.9), Inches(1.1))
    tf_m = tb_meta.text_frame
    tf_m.margin_left = tf_m.margin_top = tf_m.margin_right = tf_m.margin_bottom = 0

    pm1 = tf_m.paragraphs[0]
    pm1.text = "ACADEMIC LAB SUBMISSION DETAILS"
    pm1.font.name = FONT_MAIN
    pm1.font.size = Pt(10)
    pm1.font.bold = True
    pm1.font.color.rgb = C_BLUE_ACCENT

    pm2 = tf_m.add_paragraph()
    pm2.text = "Course / Subject : Object-Oriented Programming with C++ (C++17)\nDepartment       : Computer Science & Engineering"
    pm2.font.name = FONT_MAIN
    pm2.font.size = Pt(11)
    pm2.font.color.rgb = C_WHITE
    pm2.space_before = Pt(4)

    add_footer(slide1, 1, is_dark=True)

    # =========================================================================
    # SLIDE 2: INTRODUCTION & AIM
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_card(slide2, 0, 0, Inches(13.333), Inches(7.5), bg_color=C_BG_LIGHT, border_color=None)
    add_header(slide2, "01 | Introduction", "Project Overview & Aim", "Automating examination conduct and evaluation using Object-Oriented C++.")

    # Aim Card (Top Banner)
    add_card(slide2, Inches(0.8), Inches(1.8), Inches(11.73), Inches(1.15), bg_color=C_BLUE_LIGHT, border_color=C_BLUE_PRIMARY)
    tb_aim = slide2.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(11.33), Inches(0.95))
    tf_aim = tb_aim.text_frame
    tf_aim.word_wrap = True
    tf_aim.margin_left = tf_aim.margin_top = tf_aim.margin_right = tf_aim.margin_bottom = 0
    pa1 = tf_aim.paragraphs[0]
    pa1.text = "AIM OF THE PROJECT:"
    pa1.font.name = FONT_MAIN
    pa1.font.size = Pt(11)
    pa1.font.bold = True
    pa1.font.color.rgb = C_BLUE_PRIMARY

    pa2 = tf_aim.add_paragraph()
    pa2.text = "To design and implement a console-based Online Examination System in C++ that loads questions dynamically from an external file, conducts an interactive quiz session, evaluates candidate answers in real time, generates an instant scorecard, and persists results to an audit file."
    pa2.font.name = FONT_MAIN
    pa2.font.size = Pt(11.5)
    pa2.font.color.rgb = C_SLATE_700
    pa2.space_before = Pt(4)

    # 2 Split Cards Below
    c_w = Inches(5.72)
    c_h = Inches(3.7)

    # Card Left: Motivation
    add_card(slide2, Inches(0.8), Inches(3.15), c_w, c_h, bg_color=C_WHITE, border_color=C_SLATE_200)
    add_card(slide2, Inches(0.8), Inches(3.15), c_w, Inches(0.06), bg_color=C_AMBER, border_color=None)
    tb_m = slide2.shapes.add_textbox(Inches(1.0), Inches(3.35), c_w - Inches(0.4), c_h - Inches(0.4))
    tf_m = tb_m.text_frame
    tf_m.word_wrap = True
    tf_m.margin_left = tf_m.margin_top = tf_m.margin_right = tf_m.margin_bottom = 0
    pm = tf_m.paragraphs[0]
    pm.text = "Why Automated Quiz System?"
    pm.font.name = FONT_MAIN
    pm.font.size = Pt(13)
    pm.font.bold = True
    pm.font.color.rgb = C_SLATE_900

    add_bullet_item(tf_m, "Manual Grading Issues", "Paper quizzes take time to evaluate and are prone to human calculation mistakes.", 11)
    add_bullet_item(tf_m, "Hardcoded Questions", "Simple procedural programs hardcode questions inside the code, requiring recompilation for changes.", 11)
    add_bullet_item(tf_m, "Instant Feedback", "Students learn immediately by knowing if their answer was right or wrong with points earned.", 11)
    add_bullet_item(tf_m, "Permanent Records", "Results must be logged with timestamps so instructors can review past candidate scores.", 11)

    # Card Right: What the Program Achieves
    add_card(slide2, Inches(6.81), Inches(3.15), c_w, c_h, bg_color=C_WHITE, border_color=C_SLATE_200)
    add_card(slide2, Inches(6.81), Inches(3.15), c_w, Inches(0.06), bg_color=C_TEAL, border_color=None)
    tb_a = slide2.shapes.add_textbox(Inches(7.01), Inches(3.35), c_w - Inches(0.4), c_h - Inches(0.4))
    tf_a = tb_a.text_frame
    tf_a.word_wrap = True
    tf_a.margin_left = tf_a.margin_top = tf_a.margin_right = tf_a.margin_bottom = 0
    pa = tf_a.paragraphs[0]
    pa.text = "What This Project Demonstrates"
    pa.font.name = FONT_MAIN
    pa.font.size = Pt(13)
    pa.font.bold = True
    pa.font.color.rgb = C_SLATE_900

    add_bullet_item(tf_a, "OOP In Practice", "Real-world usage of Classes, Objects, Inheritance, and Dynamic Polymorphism.", 11, bullet_color=C_TEAL)
    add_bullet_item(tf_a, "File Handling", "Seamless reading from questions.txt and appending candidate scores to results.txt.", 11, bullet_color=C_TEAL)
    add_bullet_item(tf_a, "Multiple Question Paradigms", "Handles Multiple Choice, True/False, and Fill in the Blanks under one base class.", 11, bullet_color=C_TEAL)
    add_bullet_item(tf_a, "Safe Memory Lifecycle", "Zero memory leaks using clean virtual destructors and cleanup routines.", 11, bullet_color=C_TEAL)

    add_footer(slide2, 2)

    # =========================================================================
    # SLIDE 3: SYSTEM REQUIREMENTS
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_card(slide3, 0, 0, Inches(13.333), Inches(7.5), bg_color=C_BG_LIGHT, border_color=None)
    add_header(slide3, "02 | Specifications", "Hardware & Software Requirements", "Standard environment specifications required to compile and execute the project.")

    req_cards = [
        ("SOFTWARE REQUIREMENTS", C_BLUE_PRIMARY, [
            ("Operating System", "Windows 10 / 11 or Linux (Ubuntu, Fedora, Mint)"),
            ("Programming Language", "C++ (Standard C++17)"),
            ("Compiler", "GNU Compiler Collection (g++ v11+ / MinGW-w64) or Clang"),
            ("Text Editor / IDE", "Visual Studio Code, Code::Blocks, Dev-C++, or Terminal"),
            ("Dependencies", "None (100% pure standard C++ library)")
        ]),
        ("HARDWARE REQUIREMENTS", C_TEAL, [
            ("Processor", "Intel / AMD Dual Core 1.5 GHz or higher"),
            ("Random Access Memory (RAM)", "512 MB minimum (1 GB recommended)"),
            ("Hard Disk Space", "Less than 2 MB (Source + compiled binary)"),
            ("Input Device", "Standard Keyboard"),
            ("Output Device", "Console Monitor / Command Prompt Display")
        ])
    ]

    for i, (title, col, items) in enumerate(req_cards):
        cx = Inches(0.8 + i * 5.95)
        cy = Inches(1.8)
        cw = Inches(5.78)
        ch = Inches(4.2)

        add_card(slide3, cx, cy, cw, ch, bg_color=C_WHITE, border_color=C_SLATE_200)
        add_card(slide3, cx, cy, cw, Inches(0.08), bg_color=col, border_color=None)

        tb = slide3.shapes.add_textbox(cx + Inches(0.25), cy + Inches(0.2), cw - Inches(0.5), ch - Inches(0.35))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT_MAIN
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = col

        for pfx, txt in items:
            add_bullet_item(tf, pfx, txt, 11, space_before=11, bullet_color=col)

    # Compilation Command Callout
    cmd_box = add_card(slide3, Inches(0.8), Inches(6.15), Inches(11.73), Inches(0.7), bg_color=C_SLATE_900, border_color=None)
    tb_cmd = slide3.shapes.add_textbox(Inches(1.0), Inches(6.22), Inches(11.3), Inches(0.55))
    tf_c = tb_cmd.text_frame
    p_cmd = tf_c.paragraphs[0]
    p_cmd.text = "Compilation Command:  g++ -std=c++17 quiz_system.cpp -o quiz_system.exe"
    p_cmd.font.name = FONT_CODE
    p_cmd.font.size = Pt(11)
    p_cmd.font.bold = True
    p_cmd.font.color.rgb = RGBColor(56, 189, 248)

    add_footer(slide3, 3)

    # =========================================================================
    # SLIDE 4: KEY FEATURES OF THE SYSTEM
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_card(slide4, 0, 0, Inches(13.333), Inches(7.5), bg_color=C_BG_LIGHT, border_color=None)
    add_header(slide4, "03 | Functionality", "Key Features & Capabilities", "Core functional highlights built into the online examination console engine.")

    features = [
        ("1. Multiple Question Paradigms", "Supports Multiple Choice (4 choices), True/False, and Fill-in-the-Blank questions within a single test session.", C_BLUE_PRIMARY),
        ("2. Real-Time Feedback & Scoring", "Validates candidate answers immediately, awards specific points, and reveals correct answers on mistakes.", C_TEAL),
        ("3. Formatted Scorecard & Grades", "Generates clean ASCII scorecard displaying Roll No, Name, Score, Percentage, and Performance Grade (A+ to F).", C_PURPLE),
        ("4. File-Driven Persistence", "Reads question bank from questions.txt and automatically appends candidate score logs with timestamps to results.txt.", C_AMBER),
        ("5. Admin Question Authoring", "Instructors can interactively add new questions (MCQ, TF, FIB) directly through the console menu without touching code.", C_BLUE_PRIMARY),
        ("6. Built-In Fault Tolerance", "Self-heals missing files by generating a default question bank, and handles bad user inputs gracefully.", C_TEAL)
    ]

    card_w = Inches(5.72)
    card_h = Inches(1.45)
    for i, (f_title, f_desc, f_col) in enumerate(features):
        col_idx = i % 2
        row_idx = i // 2
        fx = Inches(0.8 + col_idx * 6.01)
        fy = Inches(1.8 + row_idx * 1.65)

        add_card(slide4, fx, fy, card_w, card_h, bg_color=C_WHITE, border_color=C_SLATE_200)
        add_card(slide4, fx, fy, Inches(0.1), card_h, bg_color=f_col, border_color=None)

        tb = slide4.shapes.add_textbox(fx + Inches(0.25), fy + Inches(0.15), card_w - Inches(0.4), card_h - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p = tf.paragraphs[0]
        p.text = f_title
        p.font.name = FONT_MAIN
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = C_SLATE_900

        p2 = tf.add_paragraph()
        p2.text = f_desc
        p2.font.name = FONT_MAIN
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = C_SLATE_600
        p2.space_before = Pt(4)

    add_footer(slide4, 4)

    # =========================================================================
    # SLIDE 5: OOP CONCEPTS IMPLEMENTED
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_card(slide5, 0, 0, Inches(13.333), Inches(7.5), bg_color=C_BG_LIGHT, border_color=None)
    add_header(slide5, "04 | OOP Concepts", "Core Object-Oriented Principles Used", "Concrete mapping of textbook OOP concepts to the project's source code.")

    oop_pillars = [
        ("DATA ABSTRACTION", C_BLUE_PRIMARY, [
            ("Concept", "Hiding internal implementation details and showing only essential features."),
            ("In Our Code", "Abstract base class Question defines pure virtual functions: display() = 0 and checkAnswer() = 0.")
        ]),
        ("DATA ENCAPSULATION", C_TEAL, [
            ("Concept", "Binding data members and member functions together while restricting direct access."),
            ("In Our Code", "Sensitive fields (score, rollNo, correctOption) are private/protected; accessed only via getters.")
        ]),
        ("HIERARCHICAL INHERITANCE", C_PURPLE, [
            ("Concept", "Deriving multiple specialized classes from a single common base class."),
            ("In Our Code", "MCQQuestion, TrueFalseQuestion, and FillInBlankQuestion all inherit publicly from Question.")
        ]),
        ("DYNAMIC POLYMORPHISM", C_AMBER, [
            ("Concept", "Resolving function calls at runtime based on the actual object type using virtual tables."),
            ("In Our Code", "QuizManager stores vector<Question*>; calls to q->display() invoke exact derived logic.")
        ])
    ]

    cw5 = Inches(5.72)
    ch5 = Inches(2.35)
    for i, (title, col, items) in enumerate(oop_pillars):
        col_idx = i % 2
        row_idx = i // 2
        px = Inches(0.8 + col_idx * 6.01)
        py = Inches(1.8 + row_idx * 2.55)

        add_card(slide5, px, py, cw5, ch5, bg_color=C_WHITE, border_color=C_SLATE_200)
        add_card(slide5, px, py, cw5, Inches(0.06), bg_color=col, border_color=None)

        tb = slide5.shapes.add_textbox(px + Inches(0.2), py + Inches(0.12), cw5 - Inches(0.4), ch5 - Inches(0.24))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT_MAIN
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = col

        for pfx, txt in items:
            add_bullet_item(tf, pfx, txt, 10.5, space_before=7, bullet_color=col)

    add_footer(slide5, 5)

    # =========================================================================
    # SLIDE 6: CLASS DESIGN & UML HIERARCHY
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_card(slide6, 0, 0, Inches(13.333), Inches(7.5), bg_color=C_BG_LIGHT, border_color=None)
    add_header(slide6, "05 | Class Design", "UML Class Architecture & Hierarchy", "Structural layout of classes, inheritance relationships, and manager components.")

    # Base Class Card (Top Center)
    bc_w = Inches(5.4)
    bc_h = Inches(2.05)
    bc_x = Inches(3.96)
    bc_y = Inches(1.8)

    add_card(slide6, bc_x, bc_y, bc_w, bc_h, bg_color=C_WHITE, border_color=C_BLUE_PRIMARY, border_width=1.5)
    add_card(slide6, bc_x, bc_y, bc_w, Inches(0.35), bg_color=C_BLUE_LIGHT, border_color=C_BLUE_PRIMARY)
    tb_bch = slide6.shapes.add_textbox(bc_x, bc_y, bc_w, Inches(0.35))
    tf_bch = tb_bch.text_frame
    pb = tf_bch.paragraphs[0]
    pb.alignment = PP_ALIGN.CENTER
    pb.text = "<<abstract>> Question"
    pb.font.name = FONT_MAIN
    pb.font.size = Pt(11)
    pb.font.bold = True
    pb.font.color.rgb = C_BLUE_PRIMARY

    tb_bcd = slide6.shapes.add_textbox(bc_x + Inches(0.15), bc_y + Inches(0.38), bc_w - Inches(0.3), Inches(1.6))
    tf_bcd = tb_bcd.text_frame
    p_b1 = tf_bcd.paragraphs[0]
    p_b1.text = "# id: int  |  # questionText: string  |  # points: int"
    p_b1.font.name = FONT_CODE
    p_b1.font.size = Pt(9.5)
    p_b1.font.color.rgb = C_SLATE_700

    p_b2 = tf_bcd.add_paragraph()
    p_b2.text = "+ virtual void display() const = 0;\n+ virtual bool checkAnswer(string userAns) const = 0;\n+ virtual string getType() const = 0;\n+ virtual ~Question() {} // Virtual Destructor"
    p_b2.font.name = FONT_CODE
    p_b2.font.size = Pt(9)
    p_b2.font.color.rgb = C_BLUE_PRIMARY
    p_b2.space_before = Pt(4)

    # 3 Derived Children Below
    child_w = Inches(3.68)
    child_h = Inches(1.85)
    children = [
        ("MCQQuestion", C_BLUE_PRIMARY, "- options: vector<string>\n- correctOption: char", "+ display() override\n+ checkAnswer() override\n+ serialize() override"),
        ("TrueFalseQuestion", C_TEAL, "- correctAnswer: bool", "+ display() override\n+ checkAnswer() override\n+ serialize() override"),
        ("FillInBlankQuestion", C_PURPLE, "- expectedAnswer: string", "+ display() override\n+ checkAnswer() override\n+ serialize() override")
    ]

    for i, (cname, col, attrs, meths) in enumerate(children):
        cx = Inches(0.8 + i * 4.02)
        cy = Inches(4.15)
        add_card(slide6, cx, cy, child_w, child_h, bg_color=C_WHITE, border_color=col, border_width=1.5)
        add_card(slide6, cx, cy, child_w, Inches(0.32), bg_color=col, border_color=None)

        tb_ch = slide6.shapes.add_textbox(cx, cy, child_w, Inches(0.32))
        tf_ch = tb_ch.text_frame
        p_ch = tf_ch.paragraphs[0]
        p_ch.alignment = PP_ALIGN.CENTER
        p_ch.text = cname
        p_ch.font.name = FONT_MAIN
        p_ch.font.size = Pt(10.5)
        p_ch.font.bold = True
        p_ch.font.color.rgb = C_WHITE

        tb_cd = slide6.shapes.add_textbox(cx + Inches(0.12), cy + Inches(0.36), child_w - Inches(0.24), Inches(1.4))
        tf_cd = tb_cd.text_frame
        p_a = tf_cd.paragraphs[0]
        p_a.text = attrs
        p_a.font.name = FONT_CODE
        p_a.font.size = Pt(9)
        p_a.font.color.rgb = C_SLATE_700

        p_me = tf_cd.add_paragraph()
        p_me.text = meths
        p_me.font.name = FONT_CODE
        p_me.font.size = Pt(8.5)
        p_me.font.color.rgb = col
        p_me.space_before = Pt(4)

    # Controller Bar
    note_box = add_card(slide6, Inches(0.8), Inches(6.2), Inches(11.73), Inches(0.65), bg_color=C_WHITE, border_color=C_SLATE_200)
    tb_n = slide6.shapes.add_textbox(Inches(1.0), Inches(6.25), Inches(11.3), Inches(0.55))
    tf_n = tb_n.text_frame
    pn = tf_n.paragraphs[0]
    pn.text = "Supporting Classes: User (candidate name, roll no, score & operator<<) | QuizManager (handles vector<Question*> and file I/O)"
    pn.font.name = FONT_MAIN
    pn.font.size = Pt(10)
    pn.font.bold = True
    pn.font.color.rgb = C_SLATE_900

    add_footer(slide6, 6)

    # =========================================================================
    # SLIDE 7: QUESTION CLASSES & INHERITANCE
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    add_card(slide7, 0, 0, Inches(13.333), Inches(7.5), bg_color=C_BG_LIGHT, border_color=None)
    add_header(slide7, "06 | Class Details", "Specialized Question Implementations", "How derived classes customize display and validation logic.")

    q_cards = [
        ("Multiple Choice (MCQ)", C_BLUE_PRIMARY, "MCQQuestion", [
            ("Purpose", "Presents a question statement along with 4 distinct options [A], [B], [C], [D]."),
            ("Data Fields", "vector<string> options for choices, char correctOption for answer key."),
            ("Validation", "Takes candidate input (e.g., 'A'), converts to uppercase, and compares with correctOption."),
            ("Serialization", "Saved in questions.txt as: MCQ|Question|OptA|OptB|OptC|OptD|Ans|Points")
        ]),
        ("True / False (TF)", C_TEAL, "TrueFalseQuestion", [
            ("Purpose", "Presents a factual statement for true or false verification."),
            ("Data Fields", "bool correctAnswer storing true or false."),
            ("Validation", "Flexible input check: accepts 'T', 'True', '1' as true, and 'F', 'False', '0' as false."),
            ("Serialization", "Saved in questions.txt as: TF|Question|True/False|Points")
        ]),
        ("Fill in the Blanks (FIB)", C_PURPLE, "FillInBlankQuestion", [
            ("Purpose", "Presents a statement with a missing keyword or phrase."),
            ("Data Fields", "string expectedAnswer storing target keyword."),
            ("Validation", "Trims leading/trailing whitespace and performs case-insensitive lexical comparison."),
            ("Serialization", "Saved in questions.txt as: FIB|Question|ExpectedAnswer|Points")
        ])
    ]

    for i, (title, col, cname, items) in enumerate(q_cards):
        cx = Inches(0.8 + i * 4.02)
        cy = Inches(1.8)
        cw = Inches(3.68)
        ch = Inches(5.0)

        add_card(slide7, cx, cy, cw, ch, bg_color=C_WHITE, border_color=C_SLATE_200)
        add_card(slide7, cx, cy, cw, Inches(0.08), bg_color=col, border_color=None)

        tb = slide7.shapes.add_textbox(cx + Inches(0.2), cy + Inches(0.2), cw - Inches(0.4), ch - Inches(0.35))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT_MAIN
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = col

        p_c = tf.add_paragraph()
        p_c.text = f"Class: {cname}"
        p_c.font.name = FONT_CODE
        p_c.font.size = Pt(9.5)
        p_c.font.color.rgb = C_SLATE_600
        p_c.space_before = Pt(2)

        for pfx, txt in items:
            add_bullet_item(tf, pfx, txt, 10.5, space_before=11, bullet_color=col)

    add_footer(slide7, 7)

    # =========================================================================
    # SLIDE 8: DYNAMIC POLYMORPHISM IN ACTION
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    add_card(slide8, 0, 0, Inches(13.333), Inches(7.5), bg_color=C_BG_LIGHT, border_color=None)
    add_header(slide8, "07 | Polymorphism", "Dynamic Polymorphism & Virtual Functions", "Executing derived class behavior through base class pointers at runtime.")

    # Left: Explanation Card
    cw8 = Inches(5.72)
    ch8 = Inches(5.0)
    add_card(slide8, Inches(0.8), Inches(1.8), cw8, ch8, bg_color=C_WHITE, border_color=C_SLATE_200)
    add_card(slide8, Inches(0.8), Inches(1.8), cw8, Inches(0.08), bg_color=C_BLUE_PRIMARY, border_color=None)

    tb_v = slide8.shapes.add_textbox(Inches(1.05), Inches(2.05), Inches(5.2), Inches(4.5))
    tf_v = tb_v.text_frame
    tf_v.word_wrap = True
    tf_v.margin_left = tf_v.margin_top = tf_v.margin_right = tf_v.margin_bottom = 0

    pv_t = tf_v.paragraphs[0]
    pv_t.text = "How Dynamic Dispatch Works"
    pv_t.font.name = FONT_MAIN
    pv_t.font.size = Pt(14)
    pv_t.font.bold = True
    pv_t.font.color.rgb = C_BLUE_PRIMARY

    add_bullet_item(tf_v, "Base Pointer Vector", "QuizManager stores vector<Question*> questions. The vector holds pointers to MCQ, True/False, and FillBlank objects together.", 11, space_before=12)
    add_bullet_item(tf_v, "Virtual Table (vtable)", "Because display() is marked virtual, C++ looks up the memory address of the actual derived function at runtime.", 11, space_before=11)
    add_bullet_item(tf_v, "No Conditional Branching", "We don't need if(type == MCQ) or switch-case statements. A single loop executes any question type cleanly.", 11, space_before=11)
    add_bullet_item(tf_v, "Extensibility", "If we add a new question type tomorrow (e.g. MatchQuestion), the quiz loop does not need any change!", 11, space_before=11)

    # Right: Code Box
    add_card(slide8, Inches(6.81), Inches(1.8), cw8, ch8, bg_color=C_SLATE_900, border_color=None)
    add_card(slide8, Inches(6.81), Inches(1.8), cw8, Inches(0.08), bg_color=C_BLUE_ACCENT, border_color=None)

    tb_code = slide8.shapes.add_textbox(Inches(7.05), Inches(2.05), Inches(5.2), Inches(4.5))
    tf_code = tb_code.text_frame
    tf_code.word_wrap = True
    tf_code.margin_left = tf_code.margin_top = tf_code.margin_right = tf_code.margin_bottom = 0

    pc_t = tf_code.paragraphs[0]
    pc_t.text = "Code Implementation (quiz_system.cpp)"
    pc_t.font.name = FONT_MAIN
    pc_t.font.size = Pt(14)
    pc_t.font.bold = True
    pc_t.font.color.rgb = C_BLUE_ACCENT

    code_snippet = (
        "// 1. Vector holding polymorphic base pointers\n"
        "vector<Question*> questions;\n\n"
        "// 2. Uniform loop executing dynamic methods\n"
        "for (Question* q : questions) {\n"
        "    // Dynamically calls MCQ::display() or\n"
        "    // TrueFalse::display() via vtable\n"
        "    q->display();\n\n"
        "    cout << \" Your Answer >> \";\n"
        "    string userAns;\n"
        "    getline(cin, userAns);\n\n"
        "    // Dynamically evaluates derived answer logic\n"
        "    if (q->checkAnswer(userAns)) {\n"
        "        student.addScore(q->getPoints());\n"
        "    }\n"
        "}\n\n"
        "// 3. Virtual destructor ensures safe cleanup\n"
        "for (Question* q : questions) delete q;"
    )

    p_body = tf_code.add_paragraph()
    p_body.text = code_snippet
    p_body.font.name = FONT_CODE
    p_body.font.size = Pt(9.5)
    p_body.font.color.rgb = RGBColor(226, 232, 240)
    p_body.space_before = Pt(12)

    add_footer(slide8, 8)

    # =========================================================================
    # SLIDE 9: FILE HANDLING & DATA STORAGE
    # =========================================================================
    slide9 = prs.slides.add_slide(blank_layout)
    add_card(slide9, 0, 0, Inches(13.333), Inches(7.5), bg_color=C_BG_LIGHT, border_color=None)
    add_header(slide9, "08 | File I/O", "File Handling & Data Persistence", "Reading questions from external repository and logging student results.")

    # Left: questions.txt
    add_card(slide9, Inches(0.8), Inches(1.8), cw8, ch8, bg_color=C_WHITE, border_color=C_SLATE_200)
    add_card(slide9, Inches(0.8), Inches(1.8), cw8, Inches(0.08), bg_color=C_BLUE_PRIMARY, border_color=None)

    tb_qf = slide9.shapes.add_textbox(Inches(1.05), Inches(2.05), Inches(5.2), Inches(4.5))
    tf_qf = tb_qf.text_frame
    tf_qf.word_wrap = True
    tf_qf.margin_left = tf_qf.margin_top = tf_qf.margin_right = tf_qf.margin_bottom = 0

    pq_t = tf_qf.paragraphs[0]
    pq_t.text = "1. Question Bank File (questions.txt)"
    pq_t.font.name = FONT_MAIN
    pq_t.font.size = Pt(14)
    pq_t.font.bold = True
    pq_t.font.color.rgb = C_BLUE_PRIMARY

    add_bullet_item(tf_qf, "Pipe Delimited Format", "Uses '|' as delimiter so regular commas and spaces in questions don't cause errors.", 11, space_before=10)
    add_bullet_item(tf_qf, "Line by Line Parsing", "Uses ifstream and getline() with stringstream to extract tokens.", 11, space_before=8)
    add_bullet_item(tf_qf, "Auto File Creation", "If questions.txt is missing on startup, the system automatically writes a default 10-question quiz.", 11, space_before=8)

    sample_q = (
        "# Sample records in questions.txt:\n"
        "MCQ|Size of empty class?|1B|0B|4B|Undef|A|2\n"
        "TF|Constructor can be virtual.|False|1\n"
        "FIB|Binding data & methods is called _____|encapsulation|2"
    )
    pq_c = tf_qf.add_paragraph()
    pq_c.text = sample_q
    pq_c.font.name = FONT_CODE
    pq_c.font.size = Pt(9.5)
    pq_c.font.color.rgb = C_SLATE_700
    pq_c.space_before = Pt(12)

    # Right: results.txt
    add_card(slide9, Inches(6.81), Inches(1.8), cw8, ch8, bg_color=C_WHITE, border_color=C_SLATE_200)
    add_card(slide9, Inches(6.81), Inches(1.8), cw8, Inches(0.08), bg_color=C_TEAL, border_color=None)

    tb_rf = slide9.shapes.add_textbox(Inches(7.05), Inches(2.05), Inches(5.2), Inches(4.5))
    tf_rf = tb_rf.text_frame
    tf_rf.word_wrap = True
    tf_rf.margin_left = tf_rf.margin_top = tf_rf.margin_right = tf_rf.margin_bottom = 0

    pr_t = tf_rf.paragraphs[0]
    pr_t.text = "2. Examination Score Log (results.txt)"
    pr_t.font.name = FONT_MAIN
    pr_t.font.size = Pt(14)
    pr_t.font.bold = True
    pr_t.font.color.rgb = C_TEAL

    add_bullet_item(tf_rf, "Append Mode (ios::app)", "Past test scores are never overwritten; every candidate test attempt is preserved.", 11, space_before=10, bullet_color=C_TEAL)
    add_bullet_item(tf_rf, "Date & Timestamp Log", "Current system date and time is recorded for every completed exam session.", 11, space_before=8, bullet_color=C_TEAL)
    add_bullet_item(tf_rf, "View Leaderboard Option", "The console menu provides Option 4 to view all past records formatted as a neat table.", 11, space_before=8, bullet_color=C_TEAL)

    sample_r = (
        "# Sample records in results.txt:\n"
        "CS202601|Alice Smith|15/17|88.2%|A (Excellent)|2026-10-15 10:30:00\n"
        "CS202602|Bob Johnson|14/17|82.4%|A (Excellent)|2026-10-15 11:15:22\n"
        "CS202603|Charlie Brown|10/17|58.8%|D (Pass)|2026-10-16 14:05:10"
    )
    pr_c = tf_rf.add_paragraph()
    pr_c.text = sample_r
    pr_c.font.name = FONT_CODE
    pr_c.font.size = Pt(9.5)
    pr_c.font.color.rgb = C_SLATE_700
    pr_c.space_before = Pt(12)

    add_footer(slide9, 9)

    # =========================================================================
    # SLIDE 10: PROGRAM EXECUTION WORKFLOW
    # =========================================================================
    slide10 = prs.slides.add_slide(blank_layout)
    add_card(slide10, 0, 0, Inches(13.333), Inches(7.5), bg_color=C_BG_LIGHT, border_color=None)
    add_header(slide10, "09 | Workflow", "Program Execution & Lifecycle", "Step-by-step flowchart of the quiz application from startup to completion.")

    steps = [
        ("Step 1", "Program Launch & File Ingestion", "Program starts, verifies questions.txt, parses question records, and instantiates objects into vector<Question*>.", C_BLUE_PRIMARY),
        ("Step 2", "Student Registration", "Candidate enters their Full Name and Roll Number, initializing the User object.", C_BLUE_ACCENT),
        ("Step 3", "Interactive Quiz Execution", "Loop runs over questions; calls q->display(), takes user input, calls q->checkAnswer(), and awards points.", C_PURPLE),
        ("Step 4", "Scorecard & Grade Generation", "Calculates total score, percentage, and letter grade; displays formatted ASCII scorecard via operator<<.", C_TEAL),
        ("Step 5", "Persistence & Memory Deallocation", "Appends candidate score to results.txt with timestamp; QuizManager safely deletes all dynamically allocated questions.", C_AMBER)
    ]

    for i, (s_num, s_title, s_desc, s_col) in enumerate(steps):
        sy = Inches(1.8 + i * 1.0)
        sw = Inches(11.73)
        sh = Inches(0.85)

        add_card(slide10, Inches(0.8), sy, sw, sh, bg_color=C_WHITE, border_color=C_SLATE_200)

        # Step badge
        nb = add_card(slide10, Inches(1.0), sy + Inches(0.15), Inches(0.95), Inches(0.55), bg_color=s_col, border_color=None)
        tb_nb = slide10.shapes.add_textbox(Inches(1.0), sy + Inches(0.15), Inches(0.95), Inches(0.55))
        tf_nb = tb_nb.text_frame
        tf_nb.margin_top = Inches(0.12)
        p_num = tf_nb.paragraphs[0]
        p_num.alignment = PP_ALIGN.CENTER
        p_num.text = s_num
        p_num.font.name = FONT_MAIN
        p_num.font.size = Pt(11)
        p_num.font.bold = True
        p_num.font.color.rgb = C_WHITE

        # Text
        tb_st = slide10.shapes.add_textbox(Inches(2.15), sy + Inches(0.12), Inches(10.1), Inches(0.65))
        tf_st = tb_st.text_frame
        tf_st.word_wrap = True
        tf_st.margin_left = tf_st.margin_top = tf_st.margin_right = tf_st.margin_bottom = 0

        p_t = tf_st.paragraphs[0]
        p_t.text = s_title
        p_t.font.name = FONT_MAIN
        p_t.font.size = Pt(12)
        p_t.font.bold = True
        p_t.font.color.rgb = C_SLATE_900

        p_d = tf_st.add_paragraph()
        p_d.text = s_desc
        p_d.font.name = FONT_MAIN
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = C_SLATE_600
        p_d.space_before = Pt(2)

    add_footer(slide10, 10)

    # =========================================================================
    # SLIDE 11: PROGRAM OUTPUT & SAMPLE SCORECARD
    # =========================================================================
    slide11 = prs.slides.add_slide(blank_layout)
    add_card(slide11, 0, 0, Inches(13.333), Inches(7.5), bg_color=C_BG_LIGHT, border_color=None)
    add_header(slide11, "10 | Output", "Console Output & Generated Scorecard", "Terminal display during question presentation and final evaluation.")

    # Left: Question Display Mockup
    add_card(slide11, Inches(0.8), Inches(1.8), cw8, ch8, bg_color=C_SLATE_900, border_color=None)
    add_card(slide11, Inches(0.8), Inches(1.8), cw8, Inches(0.08), bg_color=C_BLUE_ACCENT, border_color=None)

    tb_u1 = slide11.shapes.add_textbox(Inches(1.05), Inches(2.05), Inches(5.2), Inches(4.5))
    tf_u1 = tb_u1.text_frame
    tf_u1.word_wrap = True
    tf_u1.margin_left = tf_u1.margin_top = tf_u1.margin_right = tf_u1.margin_bottom = 0

    pu_t = tf_u1.paragraphs[0]
    pu_t.text = "Sample Question Presentation"
    pu_t.font.name = FONT_MAIN
    pu_t.font.size = Pt(14)
    pu_t.font.bold = True
    pu_t.font.color.rgb = C_BLUE_ACCENT

    q_mock = (
        "============================================================\n"
        " [Question 1] (Multiple Choice | 2 Points)\n"
        "------------------------------------------------------------\n"
        " What is the typical size of an empty class object in C++?\n\n"
        "   [A] 1 byte\n"
        "   [B] 0 bytes\n"
        "   [C] 4 bytes\n"
        "   [D] Undefined\n"
        "============================================================\n"
        " Your Answer >> A\n"
        " >> [CORRECT!] +2 points.\n"
        "------------------------------------------------------------\n"
        " [Question 2] (True / False | 1 Point)\n"
        "------------------------------------------------------------\n"
        " A constructor in C++ can be declared as virtual.\n\n"
        "   Options: Enter [T] for True or [F] for False\n"
        "============================================================\n"
        " Your Answer >> F\n"
        " >> [CORRECT!] +1 points."
    )
    p_qm = tf_u1.add_paragraph()
    p_qm.text = q_mock
    p_qm.font.name = FONT_CODE
    p_qm.font.size = Pt(8.5)
    p_qm.font.color.rgb = RGBColor(226, 232, 240)
    p_qm.space_before = Pt(8)

    # Right: Scorecard Mockup
    add_card(slide11, Inches(6.81), Inches(1.8), cw8, ch8, bg_color=C_SLATE_900, border_color=None)
    add_card(slide11, Inches(6.81), Inches(1.8), cw8, Inches(0.08), bg_color=C_TEAL, border_color=None)

    tb_u2 = slide11.shapes.add_textbox(Inches(7.05), Inches(2.05), Inches(5.2), Inches(4.5))
    tf_u2 = tb_u2.text_frame
    tf_u2.word_wrap = True
    tf_u2.margin_left = tf_u2.margin_top = tf_u2.margin_right = tf_u2.margin_bottom = 0

    pu2_t = tf_u2.paragraphs[0]
    pu2_t.text = "Generated Scorecard (operator<<)"
    pu2_t.font.name = FONT_MAIN
    pu2_t.font.size = Pt(14)
    pu2_t.font.bold = True
    pu2_t.font.color.rgb = C_TEAL

    card_mock = (
        "+-------------------------------------------------------------+\n"
        "|                     QUIZ SCORE CARD                         |\n"
        "+-------------------------------------------------------------+\n"
        "| Student Name      : Alice Smith                             |\n"
        "| Roll Number       : CS202601                                |\n"
        "| Total Questions   : 10                                      |\n"
        "| Correct Answers   : 9                                       |\n"
        "| Final Score       : 15 / 17                                 |\n"
        "| Percentage        : 88.24%                                  |\n"
        "| Performance Grade : A  (Excellent)                         |\n"
        "+-------------------------------------------------------------+\n"
        " [System] Result recorded successfully in: results.txt"
    )
    p_cm = tf_u2.add_paragraph()
    p_cm.text = card_mock
    p_cm.font.name = FONT_CODE
    p_cm.font.size = Pt(9)
    p_cm.font.color.rgb = RGBColor(226, 232, 240)
    p_cm.space_before = Pt(12)

    add_footer(slide11, 11)

    # =========================================================================
    # SLIDE 12: VIVA VOCE QUESTIONS & ANSWERS
    # =========================================================================
    slide12 = prs.slides.add_slide(blank_layout)
    add_card(slide12, 0, 0, Inches(13.333), Inches(7.5), bg_color=C_BG_LIGHT, border_color=None)
    add_header(slide12, "11 | Viva Voce", "Top Viva Questions & Answers", "Key technical questions examiners frequently ask during project defense.")

    viva_qa = [
        ("Q1: Why declare display() as pure virtual (= 0) in Question?",
         "Ans: Because MCQ, True/False, and Fill-in-the-Blank require completely different display formatting. A generic display cannot exist. Declaring it pure virtual forces every derived class to provide its own implementation and prevents instantiating the incomplete base class."),
        ("Q2: Why is the destructor virtual ~Question() in the base class?",
         "Ans: When we delete a derived object through a base pointer (delete q), a normal destructor only calls the base destructor. Marking it virtual guarantees the derived class destructor executes first, properly freeing all child memory and preventing memory leaks."),
        ("Q3: Why is operator<< overloaded as a friend function?",
         "Ans: In expressions like 'cout << user', the left operand is cout (std::ostream), which is a standard library class we cannot modify. Making it a friend function allows ostream on the left while granting access to private User data members."),
        ("Q4: How does dynamic polymorphism work at runtime here?",
         "Ans: Each object has an invisible pointer (*vptr) to a virtual table (vtable). When q->checkAnswer() executes, C++ looks up the memory address in the vtable to execute the correct derived method dynamically at runtime.")
    ]

    for i, (q_txt, a_txt) in enumerate(viva_qa):
        vy = Inches(1.8 + i * 1.25)
        vw = Inches(11.73)
        vh = Inches(1.15)

        add_card(slide12, Inches(0.8), vy, vw, vh, bg_color=C_WHITE, border_color=C_SLATE_200)
        add_card(slide12, Inches(0.8), vy, Inches(0.08), vh, bg_color=C_BLUE_PRIMARY, border_color=None)

        tb_v = slide12.shapes.add_textbox(Inches(1.1), vy + Inches(0.1), vw - Inches(0.4), vh - Inches(0.2))
        tf_v = tb_v.text_frame
        tf_v.word_wrap = True
        tf_v.margin_left = tf_v.margin_top = tf_v.margin_right = tf_v.margin_bottom = 0

        pv_q = tf_v.paragraphs[0]
        pv_q.text = q_txt
        pv_q.font.name = FONT_MAIN
        pv_q.font.size = Pt(11.5)
        pv_q.font.bold = True
        pv_q.font.color.rgb = C_SLATE_900

        pv_a = tf_v.add_paragraph()
        pv_a.text = a_txt
        pv_a.font.name = FONT_MAIN
        pv_a.font.size = Pt(9.5)
        pv_a.font.color.rgb = C_SLATE_600
        pv_a.space_before = Pt(3)

    add_footer(slide12, 12)

    # =========================================================================
    # SLIDE 13: CONCLUSION & LEARNING OUTCOMES
    # =========================================================================
    slide13 = prs.slides.add_slide(blank_layout)
    add_card(slide13, 0, 0, Inches(13.333), Inches(7.5), bg_color=C_BG_LIGHT, border_color=None)
    add_header(slide13, "12 | Conclusion", "Learning Outcomes & Summary", "Key technical achievements gained from implementing this project.")

    outcomes = [
        ("1. Mastery of OOP Principles", C_BLUE_PRIMARY, [
            "Learned how to apply Abstraction, Encapsulation, Inheritance, and Polymorphism in a practical application rather than just theory.",
            "Understood how dynamic method dispatch works through virtual tables (vtable and vptr)."
        ]),
        ("2. Safe Memory Management", C_TEAL, [
            "Gained hands-on experience managing heap memory with new and delete.",
            "Understood the critical importance of virtual destructors to prevent memory leaks."
        ]),
        ("3. Practical File Handling", C_PURPLE, [
            "Mastered reading and parsing formatted pipe-delimited text files using ifstream and stringstream.",
            "Learned to maintain persistent logs using append mode (ios::app) with timestamps."
        ]),
        ("4. Clean Academic Coding", C_AMBER, [
            "Organized a complete, working mini-project in a single clean source file (quiz_system.cpp) that compiles effortlessly in one command."
        ])
    ]

    for i, (otitle, ocol, items) in enumerate(outcomes):
        col_idx = i % 2
        row_idx = i // 2
        ox = Inches(0.8 + col_idx * 6.01)
        oy = Inches(1.8 + row_idx * 2.55)

        add_card(slide13, ox, oy, cw5, ch5, bg_color=C_WHITE, border_color=C_SLATE_200)
        add_card(slide13, ox, oy, cw5, Inches(0.06), bg_color=ocol, border_color=None)

        tb = slide13.shapes.add_textbox(ox + Inches(0.2), oy + Inches(0.12), cw5 - Inches(0.4), ch5 - Inches(0.24))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p = tf.paragraphs[0]
        p.text = otitle
        p.font.name = FONT_MAIN
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = ocol

        for itm in items:
            add_bullet_item(tf, "", itm, 10.5, space_before=8, bullet_color=ocol)

    add_footer(slide13, 13)

    # =========================================================================
    # SLIDE 14: THANK YOU SLIDE (Clean Modern Dark)
    # =========================================================================
    slide14 = prs.slides.add_slide(blank_layout)
    add_card(slide14, 0, 0, Inches(13.333), Inches(7.5), bg_color=C_SLATE_900, border_color=None)
    add_card(slide14, 0, 0, Inches(13.333), Inches(0.1), bg_color=C_BLUE_PRIMARY, border_color=None)

    tb_th = slide14.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.3), Inches(2.5))
    tf_th = tb_th.text_frame
    tf_th.word_wrap = True

    p_th1 = tf_th.paragraphs[0]
    p_th1.text = "THANK YOU!"
    p_th1.font.name = FONT_MAIN
    p_th1.font.size = Pt(44)
    p_th1.font.bold = True
    p_th1.font.color.rgb = C_WHITE

    p_th2 = tf_th.add_paragraph()
    p_th2.text = "Any Questions or Discussion?"
    p_th2.font.name = FONT_MAIN
    p_th2.font.size = Pt(22)
    p_th2.font.color.rgb = C_BLUE_ACCENT
    p_th2.space_before = Pt(8)

    p_th3 = tf_th.add_paragraph()
    p_th3.text = "Online Examination & Quiz System | C++17 OOP Mini Project"
    p_th3.font.name = FONT_MAIN
    p_th3.font.size = Pt(13)
    p_th3.font.color.rgb = C_SLATE_400
    p_th3.space_before = Pt(14)

    # Bottom Info Box
    add_card(slide14, Inches(1.0), Inches(5.0), Inches(11.33), Inches(1.3), bg_color=C_NAVY_DARK, border_color=C_SLATE_700)
    tb_binfo = slide14.shapes.add_textbox(Inches(1.2), Inches(5.15), Inches(10.9), Inches(1.0))
    tf_bi = tb_binfo.text_frame
    p_bi1 = tf_bi.paragraphs[0]
    p_bi1.text = "PROJECT REPOSITORY & FILES"
    p_bi1.font.name = FONT_MAIN
    p_bi1.font.size = Pt(10)
    p_bi1.font.bold = True
    p_bi1.font.color.rgb = C_BLUE_ACCENT

    p_bi2 = tf_bi.add_paragraph()
    p_bi2.text = "Source Code: quiz_system.cpp  |  Question Bank: questions.txt  |  Results Log: results.txt"
    p_bi2.font.name = FONT_MAIN
    p_bi2.font.size = Pt(11)
    p_bi2.font.color.rgb = C_WHITE
    p_bi2.space_before = Pt(4)

    add_footer(slide14, 14, is_dark=True)

    # Save presentation
    output_path = "Online_Quiz_System_Presentation.pptx"
    prs.save(output_path)
    print(f"[SUCCESS] 14-Slide Presentation generated successfully: {output_path}")

if __name__ == "__main__":
    create_presentation()
